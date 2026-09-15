import io
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from django.utils.translation import gettext as _
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable,
    Image,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

def _get_logo_path():
    dark = settings.BASE_DIR / 'static' / 'images' / 'logo-dark.png'
    light = settings.BASE_DIR / 'static' / 'images' / 'logo-light.png'
    if dark.exists():
        return dark
    if light.exists():
        return light
    return None


def _get_proportional_logo_image(logo_path, max_height=0.65 * inch, max_width=1.0 * inch):
    """Safely loads logo with exact aspect ratio preserved."""
    if not logo_path or not logo_path.exists():
        return None
    try:
        from PIL import Image as PILImage
        with PILImage.open(str(logo_path)) as pil_img:
            orig_w, orig_h = pil_img.size
        aspect = (orig_w / orig_h) if orig_h else 1.0
        target_h = max_height
        target_w = target_h * aspect
        if target_w > max_width:
            target_w = max_width
            target_h = target_w / aspect
        return Image(str(logo_path), width=target_w, height=target_h)
    except Exception:
        return None




# ===========================================================================
# PDF EXPORTS (ReportLab)
# ===========================================================================

def _get_pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#64748b'),
        spaceAfter=12,
    ))
    styles.add(ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#2563eb'),
        spaceBefore=10,
        spaceAfter=4,
    ))
    styles.add(ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        'MetaLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#475569'),
    ))
    styles.add(ParagraphStyle(
        'MetaValue',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#1e293b'),
    ))
    return styles


def export_report_pdf(report) -> io.BytesIO:
    """Generates a professional PDF document for a TaskReport in English."""
    from django.utils.translation import override

    with override('en'):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36,
        )
        styles = _get_pdf_styles()
        story = []

        # Header with Logo & Title
        header_data = []
        logo_path = _get_logo_path()
        logo_img = _get_proportional_logo_image(logo_path, max_height=0.65 * inch, max_width=0.8 * inch)
        if logo_img:
            header_data.append([logo_img, Paragraph(f"<b>AL-KHWARIZMI UNIVERSITY</b><br/><font size=8 color='#64748b'>Task Report | {timezone.now():%d.%m.%Y %H:%M}</font>", styles['BodyDark'])])
        else:
            header_data.append([Paragraph("<b>AL-KHWARIZMI UNIVERSITY</b>", styles['DocTitle']), Paragraph(f"<font size=8 color='#64748b'>Task Report | {timezone.now():%d.%m.%Y %H:%M}</font>", styles['BodyDark'])])

        header_table = Table(header_data, colWidths=[1.2 * inch, 5.8 * inch])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
        ]))
        story.append(header_table)
        story.append(Spacer(1, 10))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563eb'), spaceAfter=12))

        # Report Title & Meta
        story.append(Paragraph(report.title, styles['DocTitle']))
        if report.reporting_period:
            story.append(Paragraph(f"Reporting Period: {report.reporting_period} | Version: v{report.version}", styles['DocSubTitle']))

        # Metadata Grid Table
        task = report.task
        meta_rows = [
            [
                Paragraph("<b>Task Number:</b>", styles['MetaLabel']),
                Paragraph(task.task_number, styles['MetaValue']),
                Paragraph("<b>Department:</b>", styles['MetaLabel']),
                Paragraph(task.responsible_department.name if task.responsible_department else "—", styles['MetaValue']),
            ],
            [
                Paragraph("<b>Task Title:</b>", styles['MetaLabel']),
                Paragraph(task.title, styles['MetaValue']),
                Paragraph("<b>Author / Assignee:</b>", styles['MetaLabel']),
                Paragraph(report.author.display_name if report.author else "—", styles['MetaValue']),
            ],
            [
                Paragraph("<b>Report Status:</b>", styles['MetaLabel']),
                Paragraph(report.get_status_display(), styles['MetaValue']),
                Paragraph("<b>Created Date:</b>", styles['MetaLabel']),
                Paragraph(f"{report.created_at:%d.%m.%Y %H:%M}", styles['MetaValue']),
            ],
        ]
        meta_table = Table(meta_rows, colWidths=[1.3 * inch, 2.2 * inch, 1.3 * inch, 2.2 * inch])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 14))

        # Content Sections
        sections = [
            ("Executive Summary", report.summary),
            ("Completed Work & Activities", report.completed_work),
            ("Key Results & Deliverables", report.results),
            ("Identified Challenges & Problems", report.problems),
            ("Recommendations & Next Steps", report.recommendations),
            ("Conclusion", report.conclusion),
        ]

        for title, content in sections:
            if content and content.strip():
                story.append(Paragraph(title, styles['SectionHeader']))
                story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#cbd5e1'), spaceAfter=6))
                # Format lines with breaks
                for para in content.strip().split('\n'):
                    if para.strip():
                        story.append(Paragraph(para.strip(), styles['BodyDark']))
                story.append(Spacer(1, 8))

        doc.build(story)
        buffer.seek(0)
        return buffer


def export_task_pdf(task) -> io.BytesIO:
    """Generates a complete Task Completion Summary PDF in English."""
    from django.utils.translation import override

    with override('en'):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36,
        )
        styles = _get_pdf_styles()
        story = []

        # Header
        logo_path = _get_logo_path()
        logo_img = _get_proportional_logo_image(logo_path, max_height=0.65 * inch, max_width=0.8 * inch)
        if logo_img:
            header_table = Table([[logo_img, Paragraph(f"<b>AL-KHWARIZMI UNIVERSITY</b><br/><font size=8 color='#64748b'>Task Summary Report | {timezone.now():%d.%m.%Y %H:%M}</font>", styles['BodyDark'])]], colWidths=[1.2 * inch, 5.8 * inch])
        else:
            header_table = Table([[Paragraph("<b>AL-KHWARIZMI UNIVERSITY</b>", styles['DocTitle']), Paragraph(f"<font size=8 color='#64748b'>Task Summary Report | {timezone.now():%d.%m.%Y %H:%M}</font>", styles['BodyDark'])]], colWidths=[1.2 * inch, 5.8 * inch])

        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
        ]))
        story.append(header_table)
        story.append(Spacer(1, 10))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563eb'), spaceAfter=12))

    story.append(Paragraph(f"{task.task_number}: {task.title}", styles['DocTitle']))
    story.append(Spacer(1, 6))

    # Task Info Table
    assignees_str = ', '.join([a.display_name for a in task.assignees]) or '—'
    task_info = [
        [Paragraph("<b>Status:</b>", styles['MetaLabel']), Paragraph(task.get_status_display(), styles['MetaValue']), Paragraph("<b>Priority:</b>", styles['MetaLabel']), Paragraph(task.get_priority_display(), styles['MetaValue'])],
        [Paragraph("<b>Department:</b>", styles['MetaLabel']), Paragraph(task.responsible_department.name if task.responsible_department else "—", styles['MetaValue']), Paragraph("<b>Complexity:</b>", styles['MetaLabel']), Paragraph(task.get_complexity_display(), styles['MetaValue'])],
        [Paragraph("<b>Deadline:</b>", styles['MetaLabel']), Paragraph(f"{task.deadline:%d.%m.%Y}" if task.deadline else "—", styles['MetaValue']), Paragraph("<b>Progress:</b>", styles['MetaLabel']), Paragraph(f"{task.progress}%", styles['MetaValue'])],
        [Paragraph("<b>Assignees:</b>", styles['MetaLabel']), Paragraph(assignees_str, styles['MetaValue']), Paragraph("<b>Completed Date:</b>", styles['MetaLabel']), Paragraph(f"{task.completed_at:%d.%m.%Y %H:%M}" if task.completed_at else "—", styles['MetaValue'])],
    ]
    t_info = Table(task_info, colWidths=[1.3 * inch, 2.2 * inch, 1.3 * inch, 2.2 * inch])
    t_info.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t_info)
    story.append(Spacer(1, 12))

    # Description
    if task.description:
        story.append(Paragraph("Task Description", styles['SectionHeader']))
        story.append(Paragraph(task.description, styles['BodyDark']))
        story.append(Spacer(1, 10))

    # Approvals & Decisions History
    approvals = task.approvals.select_related('actor').order_by('created_at')
    if approvals.exists():
        story.append(Paragraph("Approval & Review History", styles['SectionHeader']))
        appr_rows = [[Paragraph("<b>Stage</b>", styles['MetaLabel']), Paragraph("<b>Decision</b>", styles['MetaLabel']), Paragraph("<b>Reviewer</b>", styles['MetaLabel']), Paragraph("<b>Date</b>", styles['MetaLabel']), Paragraph("<b>Notes / Reason</b>", styles['MetaLabel'])]]
        for app in approvals:
            appr_rows.append([
                Paragraph(app.get_stage_display(), styles['MetaValue']),
                Paragraph(app.get_decision_display(), styles['MetaValue']),
                Paragraph(app.actor.display_name if app.actor else "System", styles['MetaValue']),
                Paragraph(f"{app.created_at:%d.%m.%Y %H:%M}", styles['MetaValue']),
                Paragraph(app.reason or "—", styles['MetaValue']),
            ])
        t_appr = Table(appr_rows, colWidths=[1.4 * inch, 1.0 * inch, 1.5 * inch, 1.2 * inch, 1.9 * inch])
        t_appr.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_appr)
        story.append(Spacer(1, 12))

    # Work Evidence Files attached
    files = task.files.filter(is_active=True).order_by('folder__name', 'original_filename')
    if files.exists():
        story.append(Paragraph("Attached Work Evidence & Files", styles['SectionHeader']))
        file_rows = [[Paragraph("<b>Folder</b>", styles['MetaLabel']), Paragraph("<b>Filename</b>", styles['MetaLabel']), Paragraph("<b>Size</b>", styles['MetaLabel']), Paragraph("<b>Uploaded By</b>", styles['MetaLabel']), Paragraph("<b>Date</b>", styles['MetaLabel'])]]
        for f in files:
            file_rows.append([
                Paragraph(f.folder.name if f.folder else "Root", styles['MetaValue']),
                Paragraph(f.original_filename, styles['MetaValue']),
                Paragraph(f.file_size_display, styles['MetaValue']),
                Paragraph(f.uploaded_by.display_name if f.uploaded_by else "—", styles['MetaValue']),
                Paragraph(f"{f.created_at:%d.%m.%Y}", styles['MetaValue']),
            ])
        t_files = Table(file_rows, colWidths=[1.3 * inch, 2.5 * inch, 0.9 * inch, 1.3 * inch, 1.0 * inch])
        t_files.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_files)

    # Phase 9: Electronic Signature & Verification Section
    signature = task.signatures.filter(status='SIGNED').select_related('signer', 'signer__position').first()
    if not signature:
        signature = task.signatures.order_by('-signed_at').select_related('signer', 'signer__position').first()

    if signature:
        story.append(Spacer(1, 14))
        sig_color = colors.HexColor('#059669' if signature.is_valid else '#d97706')
        story.append(HRFlowable(width="100%", thickness=1.5, color=sig_color, spaceAfter=8))
        story.append(Paragraph("OFFICIAL ELECTRONIC SIGNATURE & VERIFICATION", styles['SectionHeader']))

        from signatures.qr import generate_qr_image
        ver_url = f"/verify/{signature.verification_id}/"
        qr_buf = generate_qr_image(ver_url, box_size=6, border=1)
        qr_img = Image(qr_buf, width=1.1 * inch, height=1.1 * inch)

        status_text = "VALID / AUTHENTIC" if signature.is_valid else signature.get_status_display().upper()
        position_name = signature.signer.position.name if signature.signer.position else signature.get_signature_type_display()
        sig_info_text = (
            f"<b>Status:</b> {status_text}<br/>"
            f"<b>Signed By:</b> {signature.signer.display_name} ({position_name})<br/>"
            f"<b>Signed Timestamp:</b> {signature.signed_at:%d.%m.%Y %H:%M} (UTC)<br/>"
            f"<b>Verification ID:</b> <font face='Courier'>{signature.verification_id}</font><br/>"
            f"<b>Verification Path:</b> {ver_url}<br/>"
            f"<b>Canonical Record Hash (SHA-256):</b> <font face='Courier' size=7>{signature.payload_hash}</font><br/>"
            f"<font size=7 color='#64748b'>Semantics: Electronic Signature of Official Task Record (Ed25519 cryptographic signature).</font>"
        )

        sig_table = Table([[Paragraph(sig_info_text, styles['BodyDark']), qr_img]], colWidths=[5.7 * inch, 1.3 * inch])
        bg_col = colors.HexColor('#f0fdf4' if signature.is_valid else '#fffbeb')
        border_col = colors.HexColor('#86efac' if signature.is_valid else '#fde68a')
        sig_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), bg_col),
            ('BOX', (0, 0), (-1, -1), 1, border_col),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(sig_table)

    doc.build(story)
    buffer.seek(0)
    return buffer


# ===========================================================================
# DOCX EXPORTS (python-docx)
# ===========================================================================

def export_report_docx(report) -> io.BytesIO:
    """Generates an editable Word document for a TaskReport."""
    doc = docx.Document()

    # Title
    heading = doc.add_heading(report.title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Meta paragraph
    task = report.task
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Task: {task.task_number} — {task.title}\n").bold = True
    p_meta.add_run(f"Author: {report.author.display_name if report.author else '—'}\n")
    p_meta.add_run(f"Department: {task.responsible_department.name if task.responsible_department else '—'}\n")
    p_meta.add_run(f"Status: {report.get_status_display()} | Version: v{report.version}\n")
    if report.reporting_period:
        p_meta.add_run(f"Reporting Period: {report.reporting_period}\n")
    p_meta.add_run(f"Generated: {timezone.now():%d.%m.%Y %H:%M}")

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(report.summary)

    doc.add_heading("Completed Work & Activities", level=1)
    doc.add_paragraph(report.completed_work)

    doc.add_heading("Key Results & Deliverables", level=1)
    doc.add_paragraph(report.results)

    if report.problems:
        doc.add_heading("Identified Challenges & Problems", level=1)
        doc.add_paragraph(report.problems)

    if report.recommendations:
        doc.add_heading("Recommendations & Next Steps", level=1)
        doc.add_paragraph(report.recommendations)

    if report.conclusion:
        doc.add_heading("Conclusion", level=1)
        doc.add_paragraph(report.conclusion)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def export_task_docx(task) -> io.BytesIO:
    """Generates an editable Word document for a Task summary."""
    doc = docx.Document()
    doc.add_heading(f"{task.task_number}: {task.title}", level=0)

    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Status: {task.get_status_display()}\n").bold = True
    p_meta.add_run(f"Department: {task.responsible_department.name if task.responsible_department else '—'}\n")
    p_meta.add_run(f"Priority: {task.get_priority_display()} | Complexity: {task.get_complexity_display()}\n")
    p_meta.add_run(f"Deadline: {task.deadline:%d.%m.%Y}\n" if task.deadline else "Deadline: —\n")
    p_meta.add_run(f"Progress: {task.progress}%\n")
    if task.completed_at:
        p_meta.add_run(f"Completed Date: {task.completed_at:%d.%m.%Y %H:%M}\n")

    if task.description:
        doc.add_heading("Description", level=1)
        doc.add_paragraph(task.description)

    # Approvals table
    approvals = task.approvals.select_related('actor').order_by('created_at')
    if approvals.exists():
        doc.add_heading("Approval History", level=1)
        table = doc.add_table(rows=1, cols=5)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Stage"
        hdr_cells[1].text = "Decision"
        hdr_cells[2].text = "Reviewer"
        hdr_cells[3].text = "Date"
        hdr_cells[4].text = "Reason / Notes"
        for a in approvals:
            row_cells = table.add_row().cells
            row_cells[0].text = a.get_stage_display()
            row_cells[1].text = a.get_decision_display()
            row_cells[2].text = a.actor.display_name if a.actor else "System"
            row_cells[3].text = f"{a.created_at:%d.%m.%Y %H:%M}"
            row_cells[4].text = a.reason or ""

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ===========================================================================
# XLSX EXPORTS (openpyxl)
# ===========================================================================

def export_report_xlsx(report) -> io.BytesIO:
    """Generates an Excel spreadsheet for a TaskReport."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work Report"

    header_font = Font(name='Calibri', size=13, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    label_font = Font(name='Calibri', size=11, bold=True)

    ws.merge_cells('A1:B1')
    ws['A1'] = f"TASK MANAGEMENT REPORT — {report.task.task_number}"
    ws['A1'].font = header_font
    ws['A1'].fill = header_fill
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    meta = [
        ("Report Title", report.title),
        ("Task Number", report.task.task_number),
        ("Task Title", report.task.title),
        ("Author", report.author.display_name if report.author else "—"),
        ("Department", report.task.responsible_department.name if report.task.responsible_department else "—"),
        ("Reporting Period", report.reporting_period or "—"),
        ("Status", report.get_status_display()),
        ("Version", f"v{report.version}"),
        ("Created At", f"{report.created_at:%d.%m.%Y %H:%M}"),
        ("Executive Summary", report.summary),
        ("Completed Work", report.completed_work),
        ("Key Results", report.results),
        ("Identified Problems", report.problems or "—"),
        ("Recommendations", report.recommendations or "—"),
        ("Conclusion", report.conclusion or "—"),
    ]

    row_idx = 3
    for label, val in meta:
        ws.cell(row=row_idx, column=1, value=label).font = label_font
        cell_val = ws.cell(row=row_idx, column=2, value=val)
        cell_val.alignment = Alignment(wrap_text=True)
        row_idx += 1

    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 65

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_task_xlsx(task) -> io.BytesIO:
    """Generates an Excel spreadsheet for a Task."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Task Summary"

    header_font = Font(name='Calibri', size=13, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    label_font = Font(name='Calibri', size=11, bold=True)

    ws.merge_cells('A1:D1')
    ws['A1'] = f"TASK SUMMARY — {task.task_number}: {task.title}"
    ws['A1'].font = header_font
    ws['A1'].fill = header_fill
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    meta = [
        ("Task Number", task.task_number, "Status", task.get_status_display()),
        ("Department", task.responsible_department.name if task.responsible_department else "—", "Priority", task.get_priority_display()),
        ("Deadline", f"{task.deadline:%d.%m.%Y}" if task.deadline else "—", "Complexity", task.get_complexity_display()),
        ("Progress", f"{task.progress}%", "Completed At", f"{task.completed_at:%d.%m.%Y %H:%M}" if task.completed_at else "—"),
    ]

    row_idx = 3
    for l1, v1, l2, v2 in meta:
        ws.cell(row=row_idx, column=1, value=l1).font = label_font
        ws.cell(row=row_idx, column=2, value=v1)
        ws.cell(row=row_idx, column=3, value=l2).font = label_font
        ws.cell(row=row_idx, column=4, value=v2)
        row_idx += 1

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 30

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
